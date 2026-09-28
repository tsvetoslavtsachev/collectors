"""Offline test for the SPDR GLD archive parser -- no network, a synthetic .xlsx
built in-memory with openpyxl (same shape as the real "US GLD Historical Archive"
sheet, verified 28.09.2026).

Run: python -m pytest collectors/gold/tests/test_fetch_gld.py
"""
import io
import openpyxl
import pandas as pd

from collectors.gold import fetch_gld as fg

HEADER = ["Date", "Closing Price", "Ounces of Gold per Share",
         "NAV/Share at 10:30am NYT", "Indicative Price per Share at 4:15pm NYT",
         "Mid point of bid/ask spread at 4:15pm NYT",
         "Premium/Discount of GLD Mid Point vs Indicative Value of GLD at 4:15pm NYT",
         "Daily Share Volume", "Total Ounces of Gold in the Trust",
         "Tonnes of Gold", "Total Net Asset Value in the Trust"]

ROWS = [
    ["18-Nov-2004", 44.38, 0.1, 44.2, 44.305, 44.37, 0.146, 5992000, 260000.0, 8.09, 114920000.0],
    ["19-Nov-2004", 44.78, 0.0999989, 44.559512, 44.694, 44.78, 0.192, 11655000, 1859994.06, 57.85, 828806907.2],
    ["22-Nov-2004", "US Holiday", "US Holiday", "US Holiday", "US Holiday", "US Holiday",
     "US Holiday", "US Holiday", "US Holiday", "US Holiday", "US Holiday"],
    ["23-Nov-2004", 45.05, 0.0999945, 44.812551, 44.812, 44.74, -0.16, 3139000, 2799952.98, 87.09, 1254751438.19],
]


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Disclaimer")            # the real file has this sheet first
    ws = wb.create_sheet(fg.SHEET)
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_holiday_row_is_dropped():
    df = fg.parse_archive(_xlsx_bytes())
    assert len(df) == 3
    assert "2004-11-22" not in set(df["as_of"])


def test_dates_and_types_are_clean():
    df = fg.parse_archive(_xlsx_bytes())
    assert list(df["as_of"]) == ["2004-11-18", "2004-11-19", "2004-11-23"]
    assert df["Tonnes of Gold"].dtype.kind == "f"
    assert df.loc[0, "Tonnes of Gold"] == 8.09


def test_to_records_shape():
    df = fg.parse_archive(_xlsx_bytes())
    out = fg.to_records(df)
    assert set(out) == {"etf_gld_tonnes", "etf_gld_oz"}
    assert out["etf_gld_tonnes"]["ok"] is True
    assert out["etf_gld_tonnes"]["records"][0] == {
        "as_of": "2004-11-18", "value": 8.09, "source": "SPDR GLD archive"}
    assert out["etf_gld_oz"]["records"][-1]["value"] == 2799952.98


def test_nav_and_close_kept_for_crosscheck_but_not_in_records():
    df = fg.parse_archive(_xlsx_bytes())
    assert fg.NAV_COLUMN in df.columns
    assert fg.CLOSE_COLUMN in df.columns
    out = fg.to_records(df)
    for sid in out:
        for r in out[sid]["records"]:
            assert set(r) == {"as_of", "value", "source"}
